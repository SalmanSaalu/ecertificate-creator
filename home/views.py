
import logging
import os
import re
import time
from pathlib import Path

import pandas as pd
from dotenv import load_dotenv
from pptx import Presentation

from django.conf import settings
from django.contrib import messages
from django.core.mail import EmailMessage
from django.shortcuts import render, redirect
from django.utils.crypto import get_random_string
from django.views import View

from .models import beforefetch
from .convter import ppt2pdf, export_pdf
from home.models import statusdetails

import base64
from email.message import EmailMessage

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build


load_dotenv()
logger = logging.getLogger(__name__)

BASE_DIR = Path(settings.BASE_DIR)
MEDIA_DIR = BASE_DIR / "media"
OUTPUT_DIR = MEDIA_DIR / "after_fetch"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
DEBUG_DIR = OUTPUT_DIR / "debug"
DEBUG_DIR.mkdir(parents=True, exist_ok=True)

def get_gmail_service():
    """
    Create a Gmail API service using a separate Gmail token.
    Does not modify the existing Google Drive token.
    """

    token_setting = os.getenv("GMAIL_TOKEN_FILE")

    if not token_setting:
        raise ValueError("GMAIL_TOKEN_FILE is not configured.")

    token_path = Path(token_setting)

    if not token_path.is_absolute():
        token_path = Path(settings.BASE_DIR) / token_path

    if not token_path.is_file():
        raise FileNotFoundError(
            f"Gmail token file not found: {token_path}"
        )

    credentials = Credentials.from_authorized_user_file(
        str(token_path),
        scopes=["https://www.googleapis.com/auth/gmail.send"],
    )

    if not credentials.refresh_token:
        raise ValueError(
            "Gmail token has no refresh token. "
            "Reauthorize the Gmail account."
        )

    if not credentials.valid:
        credentials.refresh(Request())

    return build(
        "gmail",
        "v1",
        credentials=credentials,
        cache_discovery=False,
    )


def send_certificate_email(recipient, pdf_data, attachment_name, row_number):
    """Send a certificate PDF through the Gmail API over HTTPS."""

    message = EmailMessage()
    message["To"] = recipient
    message["Subject"] = f"Your Certificate - Row {row_number}"

    message.set_content(
        "Congratulations!\n\n"
        "Please find your certificate attached to this email."
    )

    message.add_attachment(
        pdf_data,
        maintype="application",
        subtype="pdf",
        filename=attachment_name,
    )

    encoded_message = base64.urlsafe_b64encode(
        message.as_bytes()
    ).decode("utf-8")

    service = get_gmail_service()

    result = (
        service.users()
        .messages()
        .send(
            userId="me",
            body={"raw": encoded_message},
        )
        .execute(num_retries=2)
    )

    logger.warning(
        "GMAIL API ACCEPTED EMAIL | Row: %s | Recipient: %s | ID: %s",
        row_number,
        recipient,
        result.get("id"),
    )

    return result

def first(request):
    return render(request, "index.html")


def normalize_field(value):
    """Normalize a field name for comparison."""
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def get_placeholders(presentation):
    """Find placeholders such as <name>, <grade>, and <event>."""
    placeholders = []

    for slide in presentation.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue

            for match in re.findall(r"<[^<>]+>", shape.text):
                if match not in placeholders:
                    placeholders.append(match)

    return placeholders


def replace_placeholders(presentation, replacements):
    """Replace placeholders independently in the supplied presentation."""
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue

            for paragraph in shape.text_frame.paragraphs:
                original_text = paragraph.text
                updated_text = original_text

                for placeholder, value in replacements.items():
                    updated_text = updated_text.replace(
                        placeholder, str(value)
                    )

                if updated_text != original_text:
                    paragraph.text = updated_text


def build_row_replacements(row, placeholders, column_mapping):
    """
    Build replacement values from ONE CSV row.

    column_mapping maps each PowerPoint placeholder to its CSV column.
    """
    replacements = {}

    for placeholder in placeholders:
        column = column_mapping[placeholder]
        value = row[column]

        if pd.isna(value):
            value = ""

        replacements[placeholder] = str(value)

    return replacements


def find_recipient(row, email_column):
    """Get the email address from the current CSV row's email column."""
    value = row[email_column]

    if pd.isna(value):
        return None

    recipient = str(value).strip()

    if not re.fullmatch(
        r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+"
        r"@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+",
        recipient,
    ):
        return None

    return recipient


def cleanup_file(path, attempts=5):
    """Try to remove a temporary file, including after a brief Windows lock."""
    if path is None:
        return

    for attempt in range(attempts):
        try:
            if path.exists():
                path.unlink()
            return

        except PermissionError:
            if attempt < attempts - 1:
                time.sleep(1)
            else:
                logger.warning(
                    "Could not remove temporary file: %s",
                    path,
                    exc_info=True,
                )

        except OSError:
            logger.warning(
                "Could not remove temporary file: %s",
                path,
                exc_info=True,
            )
            return


def log_presentation_contents(presentation, stage, row_number):
    """Print all slide text so we can verify each certificate."""
    logger.warning(
        "========== CERTIFICATE DEBUG | ROW %s | %s ==========",
        row_number,
        stage,
    )

    for slide_number, slide in enumerate(
        presentation.slides, start=1
    ):
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    logger.warning(
                        "ROW %s | SLIDE %s | TEXT: %r",
                        row_number,
                        slide_number,
                        text,
                    )

    logger.warning("========== END CERTIFICATE DEBUG ==========")



class issueCertificate(View):

    def get(self, request):
        return render(request, "issuecertificate.html")

    def post(self, request):
        uploaded_csv = request.FILES.get("excel")
        uploaded_template = request.FILES.get("template")

        if not uploaded_csv or not uploaded_template:
            messages.error(
                request,
                "Please upload both a CSV file and a PowerPoint template.",
            )
            return render(request, "issuecertificate.html")

        if not uploaded_csv.name.lower().endswith(".csv"):
            messages.error(request, "Please upload a valid CSV file.")
            return render(request, "issuecertificate.html")

        if not uploaded_template.name.lower().endswith(".pptx"):
            messages.error(
                request,
                "Please save your PowerPoint template as .pptx "
                "and upload it again.",
            )
            return render(request, "issuecertificate.html")

        try:
            df = pd.read_csv(uploaded_csv)
            df.columns = [str(column).strip() for column in df.columns]

            if df.empty:
                messages.error(request, "The CSV file contains no data.")
                return render(request, "issuecertificate.html")

            presentation = Presentation(uploaded_template)
            placeholders = get_placeholders(presentation)

            if not placeholders:
                messages.error(
                    request,
                    "No placeholders were found. Use placeholders "
                    "such as <name>, <grade>, <event>, and <date>.",
                )
                return render(request, "issuecertificate.html")

            unique_id = get_random_string(length=32)

            uploaded_csv.seek(0)
            uploaded_template.seek(0)

            beforefetch.objects.create(
                excel_file_upload=uploaded_csv,
                template_file_upload=uploaded_template,
                uniquename=unique_id,
            )

            return render(
                request,
                "certificateparttwo.html",
                {
                    "forexcel": list(df.columns),
                    "fortemplate": placeholders,
                    "de": unique_id,
                },
            )

        except Exception:
            logger.exception("Error preparing certificate template.")
            messages.error(
                request,
                "Could not read the CSV or PowerPoint template. "
                "Check the server log for details.",
            )
            return render(request, "issuecertificate.html")


class Certificateparttwo(View):

    def get(self, request):
        return redirect("/")

    def post(self, request):
        event_id = request.POST.get("event")

        if not event_id:
            messages.error(request, "Missing certificate request ID.")
            return redirect("/")

        try:
            upload = beforefetch.objects.get(uniquename=event_id)

            with upload.excel_file_upload.open("rb") as csv_file:
                df = pd.read_csv(csv_file)

            # Clean headers without changing the actual column names.
            df.columns = [str(column).strip() for column in df.columns]

            if df.empty:
                messages.error(request, "The CSV file contains no data.")
                return redirect("/")

            # Identify the email column explicitly.
            email_columns = [
                column
                for column in df.columns
                if normalize_field(column) == "email"
            ]

            if not email_columns:
                messages.error(
                    request,
                    "Your CSV file must contain an email column.",
                )
                return redirect("/")

            email_column = email_columns[0]

            template_path = upload.template_file_upload.path

            original_presentation = Presentation(template_path)
            placeholders = get_placeholders(original_presentation)

            if not placeholders:
                messages.error(
                    request,
                    "No placeholders were found in the PowerPoint template.",
                )
                return redirect("/")

            # Match placeholders to CSV headers by name.
            normalized_columns = {}

            for column in df.columns:
                normalized = normalize_field(column)

                if normalized in normalized_columns:
                    raise ValueError(
                        f"Duplicate CSV headers after normalization: "
                        f"{column!r}"
                    )

                normalized_columns[normalized] = column

            column_mapping = {}

            for placeholder in placeholders:
                field_name = normalize_field(
                    placeholder.strip("<>")
                )

                if field_name not in normalized_columns:
                    raise ValueError(
                        f"PowerPoint placeholder {placeholder!r} "
                        f"does not match a CSV column. "
                        f"CSV columns: {list(df.columns)}"
                    )

                column_mapping[placeholder] = (
                    normalized_columns[field_name]
                )

            sent_count = 0
            failed_count = 0

            # Each CSV row creates a separate certificate and email.
            for row_number, (_, row) in enumerate(
                df.iterrows(), start=1
            ):
                recipient = find_recipient(row, email_column)

                if not recipient:
                    logger.error(
                        "Skipping row %s: invalid email address.",
                        row_number,
                    )
                    failed_count += 1
                    continue

                pptx_path = None
                pdf_path = None

                try:
                    # CRITICAL: replacements come from the current row.
                    replacements = build_row_replacements(
                        row,
                        placeholders,
                        column_mapping,
                    )

                    logger.info(
                        "ROW %s | RECIPIENT: %s | DATA: %s",
                        row_number,
                        recipient,
                        replacements,
                    )

                    # Reopen the ORIGINAL template for every recipient.
                    presentation = Presentation(template_path)

                    replace_placeholders(
                        presentation,
                        replacements,
                    )

                    # Give every generated file a unique name.
                    filename = get_random_string(24)

                    pptx_path = OUTPUT_DIR / f"{filename}.pptx"
                    pdf_path = OUTPUT_DIR / f"{filename}.pdf"

                    presentation.save(str(pptx_path))
                    # Verify the actual saved PowerPoint, not just in-memory values.
                    check_presentation = Presentation(str(pptx_path))

                    log_presentation_contents(
                        check_presentation,
                        "SAVED PPTX BEFORE GOOGLE UPLOAD",
                        row_number,
                    )

                    # Preserve a copy for manual inspection.
                    debug_pptx = DEBUG_DIR / f"row_{row_number}_certificate.pptx"
                    check_presentation.save(str(debug_pptx))

                    # Reopen the saved PPTX and check the actual content.
                    check_presentation = Presentation(str(pptx_path))

                    generated_text = "\n".join(
                        shape.text
                        for slide in check_presentation.slides
                        for shape in slide.shapes
                        if shape.has_text_frame
                    )

                    for placeholder in placeholders:
                        if placeholder in generated_text:
                            raise ValueError(
                                f"Placeholder {placeholder!r} remains "
                                f"unreplaced in row {row_number}."
                            )

                    logger.info(
                        "LOCAL PPTX verified for row %s: %s",
                        row_number,
                        pptx_path.name,
                    )

                    # Upload THIS recipient's generated PPTX.
                    presentation_id = ppt2pdf(
                        str(pptx_path),
                        pptx_path.name,
                        folder_id=os.getenv("GOOGLE_DRIVE_FOLDER_ID"),
                    )

                    logger.info(
                        "ROW %s | GOOGLE SLIDES ID: %s",
                        row_number,
                        presentation_id,
                    )

                    # Export the same uploaded presentation ID.
                    pdf_data = export_pdf(presentation_id)
                    logger.warning(
                    "ROW %s | Recipient: %s | Google Slides ID: %s | "
                    "PDF size: %s bytes | PDF signature: %r",
                    row_number,
                    recipient,
                    presentation_id,
                    len(pdf_data),
                    pdf_data[:8],
                    )

                    # Preserve the exact PDF that will be attached to the email.
                    debug_pdf = DEBUG_DIR / f"row_{row_number}_certificate.pdf"
                    debug_pdf.write_bytes(pdf_data)

                    if not pdf_data.startswith(b"%PDF-"):
                        raise ValueError(
                            "Google Drive returned invalid PDF data."
                        )

                    pdf_path.write_bytes(pdf_data)

                    # Send the generated PDF to THIS row's recipient.
                    #
                    attachment_name = f"certificate_row_{row_number}.pdf"

                    logger.warning(
                        "SENDING EMAIL VIA GMAIL API | Row: %s | Recipient: %s",
                        row_number,
                        recipient,
                    )

                    send_certificate_email(
                        recipient=recipient,
                        pdf_data=bytes(pdf_data),
                        attachment_name=attachment_name,
                        row_number=row_number,
                    )

                    statusdetails.objects.create(
                        email=recipient,
                        statusnamecheck=request.user.id,
                        send=1,
                    )


                    logger.warning(
                        "EMAIL SENT | Row: %s | Recipient: %s",
                        row_number,
                        recipient,
                    )

                    #

                    # if email_result != 1:
                    #     raise RuntimeError(
                    #         "Django did not confirm email submission."
                    #     )

                    statusdetails.objects.create(
                        email=recipient,
                        statusnamecheck=request.user.id,
                        send=1,
                    )

                    sent_count += 1

                    logger.info(
                        "Email submitted for row %s to %s",
                        row_number,
                        recipient,
                    )

                except Exception:
                    logger.exception(
                        "Failed to generate/send certificate "
                        "for row %s to %s",
                        row_number,
                        recipient,
                    )

                    failed_count += 1

                    try:
                        statusdetails.objects.create(
                            email=recipient,
                            statusnamecheck=request.user.id,
                            send=0,
                        )
                    except Exception:
                        logger.exception(
                            "Could not save failure status for %s",
                            recipient,
                        )

                finally:
                    cleanup_file(pptx_path)
                    cleanup_file(pdf_path)

            messages.success(
                request,
                f"Certificate processing finished. "
                f"Emails sent: {sent_count}. "
                f"Failed or skipped: {failed_count}.",
            )

        except beforefetch.DoesNotExist:
            messages.error(
                request,
                "The certificate request was not found. Please start again.",
            )

        except Exception:
            logger.exception("Certificate processing failed.")
            messages.error(
                request,
                "Certificate processing failed. Check the server logs "
                "for details.",
            )

        return redirect("/")


def status(request):
    user_id = request.user.id

    records = statusdetails.objects.filter(
        statusnamecheck=user_id
    ).order_by("-dateadded")

    return render(
        request,
        "status.html",
        {"statusdetails": records},
    )
